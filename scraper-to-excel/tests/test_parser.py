import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook  # noqa: E402

from scraper import Book, parse_page, write_csv, write_xlsx  # noqa: E402

FIXTURE = Path(__file__).parent / "fixtures" / "page-1.html"
PAGE_URL = "https://books.toscrape.com/catalogue/page-1.html"


def test_parse_page_extracts_20_books() -> None:
    books, _ = parse_page(FIXTURE.read_text(encoding="utf-8"), PAGE_URL)
    assert len(books) == 20
    assert books[0] == Book(
        title="A Light in the Attic",
        price_gbp=51.77,
        rating=3,
        in_stock=True,
        url="https://books.toscrape.com/catalogue/a-light-in-the-attic_1000/index.html",
    )
    assert all(1 <= b.rating <= 5 and b.price_gbp > 0 for b in books)


def test_parse_page_next_link() -> None:
    _, next_url = parse_page(FIXTURE.read_text(encoding="utf-8"), PAGE_URL)
    assert next_url == "https://books.toscrape.com/catalogue/page-2.html"


def test_last_page_has_no_next() -> None:
    _, next_url = parse_page("<html><ul class='pager'></ul></html>", PAGE_URL)
    assert next_url is None


def test_exports_roundtrip(tmp_path: Path) -> None:
    books, _ = parse_page(FIXTURE.read_text(encoding="utf-8"), PAGE_URL)
    write_csv(books, tmp_path / "b.csv")
    write_xlsx(books, tmp_path / "b.xlsx")
    assert len((tmp_path / "b.csv").read_text(encoding="utf-8").splitlines()) == 21
    ws = load_workbook(tmp_path / "b.xlsx").active
    assert ws.max_row == 21
    assert ws["A2"].value == "A Light in the Attic"
