"""Scrape books.toscrape.com (a sandbox built for scraping practice) to CSV and XLSX.

Usage: python scraper.py --pages 5 --out output
"""

import argparse
import csv
import logging
import time
from dataclasses import asdict, dataclass, fields
from pathlib import Path
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

BASE_URL = "https://books.toscrape.com/catalogue/page-1.html"
RATINGS = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}
REQUEST_DELAY_S = 0.5
TIMEOUT_S = 15

log = logging.getLogger("scraper")


@dataclass(frozen=True)
class Book:
    title: str
    price_gbp: float
    rating: int
    in_stock: bool
    url: str


def make_session(retries: int = 3) -> requests.Session:
    retry = Retry(total=retries, backoff_factor=1.0,
                  status_forcelist=(429, 500, 502, 503, 504), allowed_methods=("GET",))
    session = requests.Session()
    session.mount("https://", HTTPAdapter(max_retries=retry))
    session.headers["User-Agent"] = "portfolio-demo-scraper/1.0"
    return session


def parse_page(html: str, page_url: str) -> tuple[list[Book], str | None]:
    """Return the books on a listing page and the absolute URL of the next page (if any)."""
    soup = BeautifulSoup(html, "html.parser")
    books = []
    for pod in soup.select("article.product_pod"):
        link = pod.select_one("h3 a")
        rating_cls = next((c for c in pod.select_one("p.star-rating")["class"] if c in RATINGS), None)
        price_text = pod.select_one("p.price_color").get_text(strip=True)
        books.append(Book(
            title=link["title"],
            price_gbp=float(price_text.lstrip("£")),
            rating=RATINGS.get(rating_cls, 0),
            in_stock="In stock" in pod.select_one("p.availability").get_text(),
            url=urljoin(page_url, link["href"]),
        ))
    next_link = soup.select_one("li.next a")
    return books, urljoin(page_url, next_link["href"]) if next_link else None


def scrape(max_pages: int, session: requests.Session) -> list[Book]:
    books: list[Book] = []
    url: str | None = BASE_URL
    for page in range(1, max_pages + 1):
        if url is None:
            break
        resp = session.get(url, timeout=TIMEOUT_S)
        resp.raise_for_status()
        resp.encoding = "utf-8"
        page_books, url = parse_page(resp.text, resp.url)
        books.extend(page_books)
        log.info("page %d: %d books (total %d)", page, len(page_books), len(books))
        time.sleep(REQUEST_DELAY_S)
    return books


def write_csv(books: list[Book], path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=[fl.name for fl in fields(Book)])
        writer.writeheader()
        writer.writerows(asdict(b) for b in books)


def write_xlsx(books: list[Book], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Books"
    headers = [fl.name for fl in fields(Book)]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for b in books:
        ws.append([b.title, b.price_gbp, b.rating, b.in_stock, b.url])
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].number_format = "£#,##0.00"
    for col, width in zip("ABCDE", (50, 12, 8, 10, 70)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--pages", type=int, default=5, help="max listing pages (20 books each)")
    parser.add_argument("--out", type=Path, default=Path("output"))
    args = parser.parse_args()

    logging.basicConfig(level=logging.INFO, format="%(message)s")
    args.out.mkdir(parents=True, exist_ok=True)
    books = scrape(args.pages, make_session())
    write_csv(books, args.out / "books.csv")
    write_xlsx(books, args.out / "books.xlsx")
    log.info("saved %d rows to %s/books.csv and books.xlsx", len(books), args.out)


if __name__ == "__main__":
    main()
